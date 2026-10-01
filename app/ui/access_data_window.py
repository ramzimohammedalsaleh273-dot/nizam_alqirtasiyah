from __future__ import annotations
from decimal import Decimal
from typing import Any
from PySide6.QtCore import Qt
from PySide6.QtPrintSupport import QPrinter, QPrintDialog
from PySide6.QtGui import QTextDocument
from PySide6.QtWidgets import (
    QApplication,QWidget,QVBoxLayout,QHBoxLayout,QFormLayout,QLineEdit,QPushButton,QLabel,
    QTableWidget,QTableWidgetItem,QHeaderView,QAbstractItemView,QMessageBox,QApplication,
    QDialog,QDialogButtonBox,QTextEdit,QDoubleSpinBox,QSpinBox,QCheckBox,
    QComboBox,QMenu,QFileDialog,QFileDialog,QFileDialog,QInputDialog
)
from sqlalchemy import text
from app.database.connection import get_session
from app.ui.theme import APP_STYLE
from openpyxl import Workbook

FIELD_LABELS={
"id":"الرقم","code":"الكود","name":"الاسم","name_ar":"اسم الصنف","name_en":"الاسم بالإنجليزية",
"sku":"رمز الصنف SKU","barcode":"الباركود","customer_code":"رقم العميل","supplier_code":"رقم المورد",
"employee_no":"رقم الموظف","phone":"الهاتف","mobile":"الجوال","email":"البريد الإلكتروني","address":"العنوان",
"tax_number":"الرقم الضريبي","commercial_number":"السجل التجاري","description":"الوصف","notes":"الملاحظات",
"is_active":"الحالة","status":"الحالة","created_at":"تاريخ الإنشاء","updated_at":"آخر تحديث",
"company_id":"الشركة","branch_id":"الفرع","warehouse_id":"المستودع","category_id":"التصنيف",
"brand_id":"العلامة التجارية","unit_id":"الوحدة","parent_id":"الأب","group_id":"المجموعة",
"department_id":"القسم","position_id":"الوظيفة","product_type":"نوع الصنف","supplier_type":"نوع المورد",
"customer_type":"نوع العميل","warehouse_type":"نوع المستودع","cost_price":"سعر التكلفة","sale_price":"سعر البيع",
"wholesale_price":"سعر الجملة","school_price":"سعر المدارس","corporate_price":"سعر الشركات","min_price":"أقل سعر",
"reorder_point":"نقطة إعادة الطلب","min_stock":"الحد الأدنى","max_stock":"الحد الأعلى","credit_limit":"حد الائتمان",
"opening_balance":"الرصيد الافتتاحي","current_balance":"الرصيد الحالي","basic_salary":"الراتب الأساسي",
"rate":"النسبة","setting_key":"مفتاح الإعداد","setting_value":"قيمة الإعداد","value_type":"نوع القيمة",
"account_code":"رقم الحساب","account_name":"اسم الحساب","account_type":"نوع الحساب","allow_posting":"يسمح بالترحيل",
"bank_id":"البنك","account_number":"رقم الحساب","iban":"الآيبان","currency_code":"العملة",
"document_no":"رقم المستند","title":"العنوان","document_type":"نوع المستند","entity_type":"نوع الكيان",
"entity_id":"معرف الكيان","file_name":"اسم الملف","file_path":"مسار الملف","message":"الرسالة",
"read_at":"تاريخ القراءة","is_read":"مقروء"
}
HIDDEN={"password_hash","session_token","token_hash","secret","private_key","xml_content"}
READONLY={"id","created_at","updated_at","read_at","last_login_at"}
TITLES={
"companies":"الشركات","branches":"الفروع","products":"المنتجات","product_categories":"التصنيفات","units":"الوحدات",
"warehouses":"المستودعات","customers":"العملاء","suppliers":"الموردون","employees":"الموظفون","users":"المستخدمون",
"roles":"الأدوار","permissions":"الصلاحيات","banks":"البنوك","bank_accounts":"الحسابات البنكية","accounts":"دليل الحسابات",
"chart_of_accounts":"دليل الحسابات التفصيلي","cash_registers":"الصناديق","tax_rates":"الضرائب","documents":"المستندات",
"notifications":"التنبيهات","system_settings":"إعدادات النظام","stocktakes":"الجرد","stock_movements":"حركات المخزون",
"audit_logs":"سجل التدقيق","audit_log":"سجل التدقيق","security_audit":"التدقيق الأمني","sync_queue":"طابور المزامنة",
"sync_devices":"أجهزة المزامنة","approval_requests":"طلبات الاعتماد","fiscal_periods":"الفترات المالية",
"payroll_periods":"فترات الرواتب","payroll_runs":"مسيرات الرواتب","assets":"الأصول","contracts":"العقود",
"printing_services":"خدمات الطباعة","printing_orders":"طلبات الطباعة","integrations":"التكاملات",
"ecommerce_orders":"طلبات التجارة الإلكترونية"
}

class RecordDialog(QDialog):
    def __init__(self,title,fields,values=None,parent=None):
        super().__init__(parent); self.setWindowTitle(title); self.setLayoutDirection(Qt.RightToLeft)
        self.setMinimumSize(620,420); self.widgets={}; values=values or {}
        root=QVBoxLayout(self); form=QFormLayout(); form.setLabelAlignment(Qt.AlignRight)
        for f in fields:
            n,t=f["name"],(f.get("type") or "").upper()
            if n in HIDDEN or n in READONLY: continue
            if "BOOL" in t or n in {"is_active","is_read","is_locked","allow_posting","is_group"}:
                w=QCheckBox(); w.setChecked(bool(values.get(n,f.get("default") not in (None,0,"0","false","False"))))
            elif n in {"description","notes","message","address"}:
                w=QTextEdit(); w.setMinimumHeight(70); w.setPlainText("" if values.get(n) is None else str(values[n]))
            elif any(x in t for x in ("REAL","NUMERIC","DOUBLE","FLOAT")):
                w=QDoubleSpinBox(); w.setRange(-999999999999,999999999999); w.setDecimals(2)
                try:w.setValue(float(values.get(n) or 0))
                except Exception:pass
            elif "INT" in t:
                w=QSpinBox(); w.setRange(-2147483648,2147483647)
                try:w.setValue(int(values.get(n) or 0))
                except Exception:pass
            else:
                w=QLineEdit(); w.setText("" if values.get(n) is None else str(values[n]))
                w.setPlaceholderText(FIELD_LABELS.get(n,n))
            self.widgets[n]=w; form.addRow(FIELD_LABELS.get(n,n),w)
        root.addLayout(form); root.addStretch()
        b=QDialogButtonBox(QDialogButtonBox.Save|QDialogButtonBox.Cancel)
        b.accepted.connect(self.accept); b.rejected.connect(self.reject); root.addWidget(b)
    def values(self):
        out={}
        for n,w in self.widgets.items():
            if isinstance(w,QCheckBox): out[n]=1 if w.isChecked() else 0
            elif isinstance(w,QTextEdit): out[n]=w.toPlainText().strip() or None
            elif isinstance(w,(QSpinBox,QDoubleSpinBox)): out[n]=w.value()
            else: out[n]=w.text().strip() or None
        return out

class AccessDataWindow(QWidget):
    def __init__(self,table_name,title=None,columns=None,editable=True,parent=None):
        super().__init__(parent); self.table_name=table_name; self.title_text=title or TITLES.get(table_name,table_name)
        self.requested_columns=columns; self.editable=editable; self.page_size=100; self.page=0; self.total=0; self.columns=[]
        self.setWindowTitle(self.title_text); self.setMinimumSize(1100,680); self.setLayoutDirection(Qt.RightToLeft); self.setStyleSheet(APP_STYLE)
        self._build(); self._schema(); self.load()
    def _build(self):
        root=QVBoxLayout(self); root.setContentsMargins(14,14,14,10)
        h=QHBoxLayout(); self.title=QLabel(self.title_text); self.title.setObjectName("SectionTitle"); h.addWidget(self.title); h.addStretch()
        self.count=QLabel(); h.addWidget(self.count); root.addLayout(h)
        s=QHBoxLayout(); s.addWidget(QLabel("بحث:")); self.search=QLineEdit(); self.search.setPlaceholderText("بحث لحظي في كل الأعمدة…"); self.search.textChanged.connect(self._changed); s.addWidget(self.search,1)
        self.filter=QComboBox(); self.filter.addItem("كل الحالات",""); self.filter.currentIndexChanged.connect(lambda *_:self.load(reset=True)); s.addWidget(self.filter); root.addLayout(s)
        a=QHBoxLayout()
        for cap,fn,obj in [("جديد",self.add,"Success"),("فتح",self.edit,"Primary"),("تعديل",self.edit,"Primary"),("حذف",self.delete,"Danger"),("تصفية",self.advanced_search,"Secondary"),("تحديث",self.load,"Secondary"),("نسخ",self.copy_selection,"Secondary"),("Excel",self.export_excel,"Secondary"),("طباعة",self.print_table,"Secondary")]:
            b=QPushButton(cap); b.setObjectName(obj); b.clicked.connect(fn); a.addWidget(b)
        a.addStretch(); root.addLayout(a)
        self.table=QTableWidget(0,0); self.table.setSortingEnabled(True); self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers); self.table.setAlternatingRowColors(True)
        self.table.setContextMenuPolicy(Qt.CustomContextMenu); self.table.customContextMenuRequested.connect(self.menu); self.table.cellDoubleClicked.connect(lambda *_:self.edit()); self.table.itemSelectionChanged.connect(self._selection_changed)
        root.addWidget(self.table,1)
        f=QHBoxLayout(); self.status=QLabel("جاهز"); f.addWidget(self.status); f.addStretch()
        for cap,fn in [("|<",self.first),("<",self.prev),(">",self.next),(">|",self.last)]:
            b=QPushButton(cap); b.setMaximumWidth(52); b.clicked.connect(fn); f.addWidget(b)
        self.page_label=QLabel(); f.addWidget(self.page_label); root.addLayout(f)
    def _schema(self):
        with get_session() as s:
            exists=s.execute(text("SELECT 1 FROM sqlite_master WHERE type='table' AND name=:n"),{"n":self.table_name}).scalar()
            if not exists:
                self.columns=[]
                self.table.setColumnCount(1)
                self.table.setHorizontalHeaderLabels(["الحالة"])
                self.status.setText(f"الجدول غير موجود: {self.table_name}")
                for b in self.findChildren(QPushButton):
                    if b.text() in {"جديد","تعديل","حذف","فتح"}: b.setEnabled(False)
                return
            rows=s.execute(text(f'PRAGMA table_info("{self.table_name}")')).mappings().all()
        self.columns=[dict(x) for x in rows if x["name"] not in HIDDEN]
        if self.requested_columns: self.columns=[x for x in self.columns if x["name"] in set(self.requested_columns)]
        self.table.setColumnCount(len(self.columns)); self.table.setHorizontalHeaderLabels([FIELD_LABELS.get(x["name"],x["name"]) for x in self.columns])
        st=next((x["name"] for x in self.columns if x["name"] in {"status","is_active","is_read"}),None)
        if st:
            with get_session() as s: vals=[r[0] for r in s.execute(text(f'SELECT DISTINCT "{st}" FROM "{self.table_name}" WHERE "{st}" IS NOT NULL ORDER BY 1')).all()]
            self.filter.clear(); self.filter.addItem("كل الحالات","")
            for v in vals:self.filter.addItem(str(v),str(v))
        else:self.filter.hide()
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents); self.table.horizontalHeader().setStretchLastSection(True)
    def _where(self):
        clauses=[]; p={}; q=self.search.text().strip()
        if q:
            cols=[x["name"] for x in self.columns if x["name"] not in {"id","created_at","updated_at"}]
            if cols: clauses.append("("+" OR ".join(f'CAST("{c}" AS TEXT) LIKE :q' for c in cols)+")"); p["q"]=f"%{q}%"
        fv=self.filter.currentData()
        if fv!="":
            st=next((x["name"] for x in self.columns if x["name"] in {"status","is_active","is_read"}),None)
            if st: clauses.append(f'CAST("{st}" AS TEXT)=:fv'); p["fv"]=fv
        return (" WHERE "+" AND ".join(clauses)) if clauses else "",p
    def _changed(self,_): self.load(reset=True)

    def focus_search(self):
        self.search.setFocus()
        self.search.selectAll()

    def clear_search(self):
        self.search.clear()
        self.filter.setCurrentIndex(0)

    def _selection_changed(self):
        self.status.setText(f"السجلات: {self.total:,}    المحدد: {1 if self.table.currentRow()>=0 else 0}")
    def load(self,*_,reset=False):
        if reset:self.page=0
        if not self.columns:return
        where,p=self._where(); cols=[x["name"] for x in self.columns]; sel=",".join(f'"{x}"' for x in cols)
        with get_session() as s:
            self.total=int(s.execute(text(f'SELECT COUNT(*) FROM "{self.table_name}"{where}'),p).scalar() or 0)
            rows=s.execute(text(f'SELECT {sel} FROM "{self.table_name}"{where} ORDER BY rowid DESC LIMIT :lim OFFSET :off'),{**p,"lim":self.page_size,"off":self.page*self.page_size}).all()
        self.table.setSortingEnabled(False); self.table.setRowCount(0)
        for row in rows:
            r=self.table.rowCount(); self.table.insertRow(r)
            for c,v in enumerate(row):
                it=QTableWidgetItem("" if v is None else str(v))
                if isinstance(v,(int,float,Decimal)):it.setTextAlignment(Qt.AlignCenter)
                self.table.setItem(r,c,it)
        self.table.setSortingEnabled(True); pages=max(1,(self.total+self.page_size-1)//self.page_size); self.page=min(self.page,pages-1)
        self.page_label.setText(f"{self.page+1} / {pages}"); self.count.setText(f"إجمالي السجلات: {self.total:,}")
        self.status.setText(f"السجلات: {self.total:,}    المحدد: {1 if self.table.currentRow()>=0 else 0}")

    def _selection_changed(self):
        self.status.setText(f"السجلات: {self.total:,}    المحدد: {len(self.table.selectionModel().selectedRows())}")

    def advanced_search(self):
        value, ok = QInputDialog.getText(self, "بحث متقدم", "ابحث في الحقول المعروضة:")
        if ok:
            self.search.setText(value)

    def copy_selection(self):
        ranges = self.table.selectedRanges()
        if not ranges:
            return
        rows = []
        for rg in ranges:
            for r in range(rg.topRow(), rg.bottomRow() + 1):
                vals = []
                for col in range(rg.leftColumn(), rg.rightColumn() + 1):
                    it = self.table.item(r, col)
                    vals.append("" if it is None else it.text())
                rows.append("\t".join(vals))
        QApplication.clipboard().setText("\n".join(rows))
        self.status.setText("تم نسخ البيانات المحددة.")

    def export_excel(self):
        path, _ = QFileDialog.getSaveFileName(self, "تصدير Excel", f"{self.title_text}.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = self.title_text[:31]
            for col, x in enumerate(self.columns, 1):
                ws.cell(1, col, FIELD_LABELS.get(x["name"], x["name"]))
            for row in range(self.table.rowCount()):
                for col in range(self.table.columnCount()):
                    it = self.table.item(row, col)
                    ws.cell(row + 2, col + 1, "" if it is None else it.text())
            wb.save(path)
            self.status.setText(f"تم التصدير: {path}")
        except Exception as e:
            QMessageBox.critical(self, "فشل التصدير", str(e))

    def print_table(self):
        try:
            from PySide6.QtPrintSupport import QPrinter, QPrintDialog
            from PySide6.QtGui import QTextDocument
            printer = QPrinter(QPrinter.HighResolution)
            dialog = QPrintDialog(printer, self)
            if dialog.exec() != QDialog.Accepted:
                return
            headers = [self.table.horizontalHeaderItem(i).text() for i in range(self.table.columnCount())]
            html = ["<html><body dir='rtl'><h2>" + self.title_text + "</h2><table border='1' cellspacing='0' cellpadding='4'><tr>"]
            html += [f"<th>{h}</th>" for h in headers]
            html.append("</tr>")
            for r in range(self.table.rowCount()):
                html.append("<tr>" + "".join(f"<td>{'' if self.table.item(r,c) is None else self.table.item(r,c).text()}</td>" for c in range(self.table.columnCount())) + "</tr>")
            html.append("</table></body></html>")
            doc = QTextDocument()
            doc.setHtml("".join(html))
            doc.print(printer)
            self.status.setText("تم إرسال الجدول إلى الطابعة.")
        except Exception as e:
            QMessageBox.critical(self, "فشل الطباعة", str(e))

    def _id(self):
        r=self.table.currentRow(); ic=next((i for i,x in enumerate(self.columns) if x["name"]=="id"),None)
        if r<0 or ic is None:return None
        try:return int(self.table.item(r,ic).text())
        except Exception:return None
    def _record(self,rid):
        with get_session() as s:return s.execute(text(f'SELECT * FROM "{self.table_name}" WHERE id=:id'),{"id":rid}).mappings().first()
    def add(self):
        if not self.editable:return
        d=RecordDialog(f"إضافة — {self.title_text}",self.columns,parent=self)
        if d.exec()!=QDialog.Accepted:return
        vals={k:v for k,v in d.values().items() if k not in READONLY and k in {x["name"] for x in self.columns}}
        if not vals:return
        try:
            with get_session() as s:
                ks=list(vals); sql=f'INSERT INTO "{self.table_name}" ({",".join(chr(34)+k+chr(34) for k in ks)}) VALUES ({",".join(":"+k for k in ks)})'
                s.execute(text(sql),vals); s.commit()
            self.load(reset=True)
        except Exception as e:QMessageBox.critical(self,"تعذر الحفظ",str(e))
    def edit(self):
        if not self.editable:return
        rid=self._id()
        if rid is None:return QMessageBox.warning(self,"تعديل","حدد سجلًا أولاً.")
        rec=self._record(rid)
        if not rec:return
        d=RecordDialog(f"تعديل — {self.title_text}",self.columns,dict(rec),self)
        if d.exec()!=QDialog.Accepted:return
        vals={k:v for k,v in d.values().items() if k not in READONLY and k in {x["name"] for x in self.columns}}
        if not vals:return
        try:
            with get_session() as s:
                sets=",".join(f'"{k}"=:{k}' for k in vals); s.execute(text(f'UPDATE "{self.table_name}" SET {sets} WHERE id=:__id'),{**vals,"__id":rid}); s.commit()
            self.load()
        except Exception as e:QMessageBox.critical(self,"تعذر التعديل",str(e))
    def delete(self):
        if not self.editable:return
        rid=self._id()
        if rid is None:return QMessageBox.warning(self,"حذف","حدد سجلًا أولاً.")
        if QMessageBox.question(self,"تأكيد الحذف","سيتم حذف السجل المحدد. هل تريد المتابعة؟")!=QMessageBox.Yes:return
        try:
            with get_session() as s:s.execute(text(f'DELETE FROM "{self.table_name}" WHERE id=:id'),{"id":rid}); s.commit()
            self.load()
        except Exception as e:QMessageBox.critical(self,"تعذر الحذف",str(e))
    def export_data(self):
        path, _ = QFileDialog.getSaveFileName(self, "تصدير البيانات", f"{self.table_name}.xlsx", "Excel (*.xlsx);;CSV (*.csv);;PDF (*.pdf)")
        if not path:
            return
        try:
            headers=[self.table.horizontalHeaderItem(i).text() for i in range(self.table.columnCount())]
            rows=[[self.table.item(r,c).text() if self.table.item(r,c) else "" for c in range(self.table.columnCount())] for r in range(self.table.rowCount())]
            if path.lower().endswith(".csv"):
                import csv
                with open(path,"w",newline="",encoding="utf-8-sig") as fh:
                    w=csv.writer(fh); w.writerow(headers); w.writerows(rows)
            elif path.lower().endswith(".pdf"):
                from reportlab.lib import colors
                from reportlab.lib.pagesizes import landscape,A4
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.ttfonts import TTFont
                doc=SimpleDocTemplate(path,pagesize=landscape(A4),rightMargin=18,leftMargin=18,topMargin=18,bottomMargin=18)
                styles=getSampleStyleSheet()
                story=[Paragraph(self.title_text,styles["Title"])]
                data=[headers]+rows
                tbl=Table(data,repeatRows=1)
                tbl.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.4,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("FONTSIZE",(0,0),(-1,-1),7),("ALIGN",(0,0),(-1,-1),"CENTER")]))
                story.append(tbl); doc.build(story)
            else:
                from openpyxl import Workbook
                wb=Workbook(); ws=wb.active; ws.title=self.title_text[:31] or "بيانات"
                ws.append(headers)
                for row in rows: ws.append(row)
                ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
                wb.save(path)
            self.status.setText(f"تم التصدير: {path}")
        except Exception as exc:
            QMessageBox.critical(self,"تعذر التصدير",str(exc))

    def print_data(self):
        try:
            printer=QPrinter(QPrinter.HighResolution)
            dialog=QPrintDialog(printer,self)
            if dialog.exec()!=QPrintDialog.Accepted:
                return
            headers=[self.table.horizontalHeaderItem(i).text() for i in range(self.table.columnCount())]
            html=["<html><meta charset='utf-8'><body dir='rtl'>",f"<h2>{self.title_text}</h2>","<table border='1' cellspacing='0' cellpadding='4' width='100%'><thead><tr>"]
            html.extend(f"<th>{h}</th>" for h in headers)
            html.append("</tr></thead><tbody>")
            for r in range(self.table.rowCount()):
                html.append("<tr>")
                for col in range(self.table.columnCount()):
                    v=self.table.item(r,col).text() if self.table.item(r,col) else ""
                    html.append(f"<td>{v}</td>")
                html.append("</tr>")
            html.append("</tbody></table></body></html>")
            doc=QTextDocument()
            doc.setHtml("".join(html))
            doc.print_(printer)
            self.status.setText("تم إرسال الجدول إلى الطابعة.")
        except Exception as exc:
            QMessageBox.critical(self,"تعذر الطباعة",str(exc))

    def menu(self,pos):
        m=QMenu(self); m.addAction("فتح / تعديل",self.edit); m.addAction("نسخ المحدد",self.copy_selected); m.addAction("تحديث",self.load)
        if self.editable:m.addSeparator();m.addAction("حذف",self.delete)
        m.exec(self.table.viewport().mapToGlobal(pos))
    def copy_selected(self):
        rows=sorted({i.row() for i in self.table.selectedIndexes()})
        if not rows:return
        QApplication.clipboard().setText("\n".join("\t".join(self.table.item(r,col).text() if self.table.item(r,col) else "" for col in range(self.table.columnCount())) for r in rows))

    def export_excel(self):
        if not self.columns:return
        path,_=QFileDialog.getSaveFileName(self,"تصدير Excel",f"{self.title_text}.xlsx","Excel (*.xlsx)")
        if not path:return
        try:
            wb=Workbook(); ws=wb.active; ws.title="البيانات"
            ws.append([FIELD_LABELS.get(x["name"],x["name"]) for x in self.columns])
            for r in range(self.table.rowCount()):
                ws.append([self.table.item(r,col).text() if self.table.item(r,col) else "" for col in range(self.table.columnCount())])
            wb.save(path); self.status.setText("تم التصدير إلى Excel.")
        except Exception as e: QMessageBox.critical(self,"فشل التصدير",str(e))

    def print_table(self):
        from PySide6.QtPrintSupport import QPrinter,QPrintDialog
        from PySide6.QtGui import QTextDocument
        if not self.columns:return
        printer=QPrinter(QPrinter.HighResolution); dlg=QPrintDialog(printer,self)
        if dlg.exec()!=QPrintDialog.Accepted:return
        headers=[FIELD_LABELS.get(x["name"],x["name"]) for x in self.columns]
        html="<html dir='rtl'><meta charset='utf-8'><h2>"+self.title_text+"</h2><table border='1' cellspacing='0' cellpadding='4'><tr>"+''.join(f"<th>{h}</th>" for h in headers)+"</tr>"
        for r in range(self.table.rowCount()):
            html+="<tr>"+''.join(f"<td>{self.table.item(r,col).text() if self.table.item(r,col) else ''}</td>" for col in range(self.table.columnCount()))+"</tr>"
        html+="</table></html>"
        doc=QTextDocument(self); doc.setHtml(html); doc.print_(printer)

    def first(self):self.page=0;self.load()
    def prev(self):self.page=max(0,self.page-1);self.load()
    def next(self):self.page=min(max(0,(self.total+self.page_size-1)//self.page_size-1),self.page+1);self.load()
    def last(self):self.page=max(0,(self.total+self.page_size-1)//self.page_size-1);self.load()
